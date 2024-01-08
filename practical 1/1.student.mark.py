import functions as fn 
from openpyxl import load_workbook
from random import randint

wb = load_workbook("studentData.xlsx")
ws = wb.active

fn.getNumberOfCourses(3)
fn.getNumberOfStudent(469)

#input courses
fn.getCourseInfomation("APP","Advanced Programming with Python")
fn.getCourseInfomation("OOP","Object Oriented Programming")
fn.getCourseInfomation("ADS","Algorithm and Data Structure")

#input students information and mark
for i in range(2,471):
    name = ws.cell(row=i,column=2).value +" "+ ws.cell(row=i,column=3).value
    dob = str(ws.cell(row=i,column=4).value)[:-9]
    id = ws.cell(row=i,column=1).value
    fn.getStudentInformation(name, id, dob)
    fn.getStudentMark(id,"OOP",randint(5,18))
    fn.getStudentMark(id,"APP",randint(5,18))
    fn.getStudentMark(id,"ADS",randint(5,18))

print("Students:\n")
fn.listStudent()
print("Courses:\n")
fn.listCourses()
print("Marks:\n")
fn.listMark("OOP")
print("--------------------\n")
fn.listMark("ADS")
print("--------------------:\n")
fn.listMark("APP")